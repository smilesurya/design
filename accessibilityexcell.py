import os
import pikepdf

from datetime import datetime
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# ============================================================
# CONFIGURATION
# ============================================================

FOLDER_PATH = r"G:\testing"

EXCEL_PATH = os.path.join(
    FOLDER_PATH,
    "Accessibility_Report.xlsx"
)


# ============================================================
# BUILD PAGE MAP
# ============================================================

def build_page_map(pdf):
    """
    Map PDF page object -> page number.
    Page numbers start from 1.
    """

    page_map = {}

    for page_number, page in enumerate(
        pdf.pages,
        start=1
    ):

        try:

            page_map[
                page.obj.objgen
            ] = page_number

        except Exception:

            pass

    return page_map


# ============================================================
# GET STRUCTURE ELEMENT PAGE NUMBER
# ============================================================

def get_struct_page_number(
    struct_elem,
    page_map,
    inherited_page=None
):
    """
    Get page number from /Pg.

    If /Pg is missing,
    use parent/inherited page.
    """

    current_page = inherited_page

    try:

        if "/Pg" in struct_elem:

            pg = struct_elem["/Pg"]

            try:

                objgen = pg.objgen

                if objgen in page_map:

                    current_page = page_map[
                        objgen
                    ]

            except Exception:

                pass

    except Exception:

        pass

    return current_page


# ============================================================
# WALK PDF STRUCTURE TREE
# ============================================================

def walk_structure(
    node,
    page_map,
    page_counts,
    total_counts,
    inherited_page=None
):

    # --------------------------------------------------------
    # ARRAY
    # --------------------------------------------------------

    if isinstance(
        node,
        pikepdf.Array
    ):

        for child in node:

            walk_structure(
                child,
                page_map,
                page_counts,
                total_counts,
                inherited_page
            )

        return


    # --------------------------------------------------------
    # Ignore MCID integers / other objects
    # --------------------------------------------------------

    if not isinstance(
        node,
        pikepdf.Dictionary
    ):
        return


    # --------------------------------------------------------
    # FIND PAGE NUMBER
    # --------------------------------------------------------

    current_page = get_struct_page_number(
        node,
        page_map,
        inherited_page
    )


    # --------------------------------------------------------
    # GET TAG NAME
    # --------------------------------------------------------

    tag_name = None

    try:

        if "/S" in node:

            tag_name = str(
                node["/S"]
            ).replace(
                "/",
                ""
            ).strip()

    except Exception:

        pass


    # --------------------------------------------------------
    # ACCESSIBILITY TAGS TO COUNT
    # --------------------------------------------------------

    wanted_tags = {
        "Note",
        "Figure",
        "Table"
    }


    # --------------------------------------------------------
    # COUNT TAG
    # --------------------------------------------------------

    if tag_name in wanted_tags:

        total_counts[
            tag_name
        ] += 1

        if current_page is not None:

            page_counts[
                current_page
            ][
                tag_name
            ] += 1


    # --------------------------------------------------------
    # WALK CHILDREN
    # --------------------------------------------------------

    try:

        if "/K" in node:

            children = node[
                "/K"
            ]

            walk_structure(
                children,
                page_map,
                page_counts,
                total_counts,
                current_page
            )

    except Exception:

        pass


# ============================================================
# ANALYZE ONE PDF
# ============================================================

def analyze_pdf(pdf_path):

    page_counts = defaultdict(
        lambda: {
            "Note": 0,
            "Figure": 0,
            "Table": 0
        }
    )

    total_counts = {
        "Note": 0,
        "Figure": 0,
        "Table": 0
    }


    print(
        f"\nReading PDF: {pdf_path}"
    )


    with pikepdf.open(
        pdf_path
    ) as pdf:

        # ----------------------------------------------------
        # PAGE COUNT
        # ----------------------------------------------------

        total_pages = len(
            pdf.pages
        )


        # ----------------------------------------------------
        # CHECK TAGGED PDF
        # ----------------------------------------------------

        if "/StructTreeRoot" not in pdf.Root:

            print(
                "WARNING: PDF does not contain StructTreeRoot."
            )

            return {
                "pages": total_pages,
                "Note": 0,
                "Figure": 0,
                "Table": 0
            }


        struct_root = pdf.Root[
            "/StructTreeRoot"
        ]


        # ----------------------------------------------------
        # BUILD PAGE MAP
        # ----------------------------------------------------

        page_map = build_page_map(
            pdf
        )


        # ----------------------------------------------------
        # WALK STRUCTURE TREE
        # ----------------------------------------------------

        if "/K" in struct_root:

            walk_structure(
                struct_root["/K"],
                page_map,
                page_counts,
                total_counts
            )


    return {
        "pages": total_pages,
        "Note": total_counts["Note"],
        "Figure": total_counts["Figure"],
        "Table": total_counts["Table"]
    }


# ============================================================
# CREATE EXCEL IF NOT AVAILABLE
# ============================================================

def create_excel():

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        "Accessibility Report"
    )


    # --------------------------------------------------------
    # HEADERS
    # --------------------------------------------------------

    headers = [
        "Date",
        "Chapter",
        "Note",
        "figure",
        "table",
        "Page count"
    ]


    worksheet.append(
        headers
    )


    # --------------------------------------------------------
    # HEADER STYLE
    # --------------------------------------------------------

    yellow_fill = PatternFill(
        fill_type="solid",
        fgColor="FFFF00"
    )

    bold_font = Font(
        bold=True
    )

    center_alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    thin_border = Border(
        left=Side(
            style="thin",
            color="000000"
        ),
        right=Side(
            style="thin",
            color="000000"
        ),
        top=Side(
            style="thin",
            color="000000"
        ),
        bottom=Side(
            style="thin",
            color="000000"
        )
    )


    for cell in worksheet[1]:

        cell.fill = yellow_fill

        cell.font = bold_font

        cell.alignment = center_alignment

        cell.border = thin_border


    # --------------------------------------------------------
    # COLUMN WIDTHS
    # --------------------------------------------------------

    worksheet.column_dimensions[
        "A"
    ].width = 15

    worksheet.column_dimensions[
        "B"
    ].width = 45

    worksheet.column_dimensions[
        "C"
    ].width = 12

    worksheet.column_dimensions[
        "D"
    ].width = 12

    worksheet.column_dimensions[
        "E"
    ].width = 12

    worksheet.column_dimensions[
        "F"
    ].width = 15


    workbook.save(
        EXCEL_PATH
    )


# ============================================================
# CHECK WHETHER CHAPTER ALREADY EXISTS
# ============================================================

def find_existing_chapter(
    worksheet,
    chapter_name
):

    for row in range(
        2,
        worksheet.max_row + 1
    ):

        value = worksheet.cell(
            row=row,
            column=2
        ).value

        if value is None:
            continue

        if str(
            value
        ).strip().lower() == chapter_name.strip().lower():

            return row

    return None


# ============================================================
# WRITE RESULT TO EXCEL
# ============================================================

def write_result_to_excel(
    chapter,
    result
):

    # --------------------------------------------------------
    # CREATE EXCEL FIRST TIME
    # --------------------------------------------------------

    if not os.path.exists(
        EXCEL_PATH
    ):

        create_excel()


    # --------------------------------------------------------
    # OPEN EXISTING EXCEL
    # --------------------------------------------------------

    workbook = load_workbook(
        EXCEL_PATH
    )

    worksheet = workbook[
        "Accessibility Report"
    ]


    # --------------------------------------------------------
    # CURRENT DATE
    # --------------------------------------------------------

    current_date = datetime.now().strftime(
        "%d-%m-%Y"
    )


    # --------------------------------------------------------
    # CHECK IF PDF IS ALREADY IN EXCEL
    # --------------------------------------------------------

    existing_row = find_existing_chapter(
        worksheet,
        chapter
    )


    # --------------------------------------------------------
    # UPDATE EXISTING RECORD
    # --------------------------------------------------------

    if existing_row:

        row = existing_row

        print(
            f"Updating existing Excel row: {row}"
        )


    # --------------------------------------------------------
    # ADD NEW RECORD
    # --------------------------------------------------------

    else:

        row = worksheet.max_row + 1

        print(
            f"Adding new Excel row: {row}"
        )


    # --------------------------------------------------------
    # WRITE DATA
    # --------------------------------------------------------

    worksheet.cell(
        row=row,
        column=1
    ).value = current_date

    worksheet.cell(
        row=row,
        column=2
    ).value = chapter

    worksheet.cell(
        row=row,
        column=3
    ).value = result[
        "Note"
    ]

    worksheet.cell(
        row=row,
        column=4
    ).value = result[
        "Figure"
    ]

    worksheet.cell(
        row=row,
        column=5
    ).value = result[
        "Table"
    ]

    worksheet.cell(
        row=row,
        column=6
    ).value = result[
        "pages"
    ]


    # --------------------------------------------------------
    # CELL ALIGNMENT + BORDER
    # --------------------------------------------------------

    center_alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    thin_border = Border(
        left=Side(
            style="thin",
            color="D9D9D9"
        ),
        right=Side(
            style="thin",
            color="D9D9D9"
        ),
        top=Side(
            style="thin",
            color="D9D9D9"
        ),
        bottom=Side(
            style="thin",
            color="D9D9D9"
        )
    )


    for column in range(
        1,
        7
    ):

        cell = worksheet.cell(
            row=row,
            column=column
        )

        cell.alignment = center_alignment

        cell.border = thin_border


    workbook.save(
        EXCEL_PATH
    )


# ============================================================
# GET ALL PDF FILES
# ============================================================

def get_pdf_files():

    pdf_files = []

    if not os.path.exists(
        FOLDER_PATH
    ):

        raise RuntimeError(
            f"Folder does not exist: {FOLDER_PATH}"
        )


    for file_name in os.listdir(
        FOLDER_PATH
    ):

        if file_name.lower().endswith(
            ".pdf"
        ):

            pdf_files.append(
                file_name
            )


    pdf_files.sort()

    return pdf_files


# ============================================================
# MAIN AUTOMATION
# ============================================================

def main():

    print(
        "=" * 70
    )

    print(
        "PDF ACCESSIBILITY TAG COUNT AUTOMATION"
    )

    print(
        "=" * 70
    )


    print(
        f"\nFolder: {FOLDER_PATH}"
    )

    print(
        f"Excel: {EXCEL_PATH}"
    )


    # --------------------------------------------------------
    # FIND PDFs
    # --------------------------------------------------------

    pdf_files = get_pdf_files()


    if not pdf_files:

        print(
            "\nNo PDF files found."
        )

        return


    print(
        f"\nPDF files found: {len(pdf_files)}"
    )


    # --------------------------------------------------------
    # PROCESS EACH PDF
    # --------------------------------------------------------

    for index, pdf_file in enumerate(
        pdf_files,
        start=1
    ):

        print(
            "\n" +
            "=" * 70
        )

        print(
            f"[{index}/{len(pdf_files)}] Processing: {pdf_file}"
        )


        pdf_path = os.path.join(
            FOLDER_PATH,
            pdf_file
        )


        try:

            # ------------------------------------------------
            # ANALYZE PDF
            # ------------------------------------------------

            result = analyze_pdf(
                pdf_path
            )


            # ------------------------------------------------
            # CHAPTER = PDF FILENAME WITHOUT .PDF
            # ------------------------------------------------

            chapter = os.path.splitext(
                pdf_file
            )[0]


            # ------------------------------------------------
            # PRINT RESULT
            # ------------------------------------------------

            print(
                f"Chapter    : {chapter}"
            )

            print(
                f"Page count : {result['pages']}"
            )

            print(
                f"Note       : {result['Note']}"
            )

            print(
                f"Figure     : {result['Figure']}"
            )

            print(
                f"Table      : {result['Table']}"
            )


            # ------------------------------------------------
            # WRITE INTO EXCEL
            # ------------------------------------------------

            write_result_to_excel(
                chapter,
                result
            )


            print(
                "Excel updated successfully."
            )


        except Exception as e:

            print(
                f"ERROR processing {pdf_file}: {e}"
            )


    # --------------------------------------------------------
    # COMPLETE
    # --------------------------------------------------------

    print(
        "\n" +
        "=" * 70
    )

    print(
        "COMPLETED"
    )

    print(
        "=" * 70
    )

    print(
        f"G:\testing\chapter01.pdf"
    )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    main()
